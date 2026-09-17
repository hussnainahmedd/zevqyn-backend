"""Document text extraction logic."""

from __future__ import annotations

import io
from uuid import UUID

import docx
import pymupdf
from fastapi import HTTPException, status

from app.core.supabase import get_admin_client
from app.models.extraction import ExtractedDocument, ExtractedUnit
from app.services.documents import get_document, BUCKET_NAME
from app.utils.text import normalize_text


MAX_PAGES = 1000
MAX_CHARS = 5_000_000

def _check_char_bound(total: int, new_length: int) -> int:
    if total + new_length > MAX_CHARS:
        raise ValueError(f"Extracted text exceeds safe limit of {MAX_CHARS} characters")
    return total + new_length

def _extract_pdf(content: bytes) -> list[ExtractedUnit]:
    """Extract text from PDF page by page."""
    units = []
    try:
        doc = pymupdf.Document(stream=content, filetype="pdf")
    except Exception as e:
        raise ValueError("Failed to parse PDF document")

    if doc.is_encrypted:
        raise ValueError("Cannot extract text from encrypted/password-protected PDF")

    if doc.page_count > MAX_PAGES:
        raise ValueError(f"Document exceeds maximum page limit of {MAX_PAGES}")

    total_chars = 0
    for i, page in enumerate(doc):
        text = page.get_text()
        norm_text = normalize_text(text)
        if norm_text:
            total_chars = _check_char_bound(total_chars, len(norm_text))
            units.append(
                ExtractedUnit(
                    index=len(units),
                    text=norm_text,
                    page_number=i + 1,
                    source_label=f"Page {i + 1}",
                )
            )

    doc.close()

    if not units:
        raise ValueError("No extractable text found in PDF (may be image-only/scanned)")

    return units


def _extract_docx(content: bytes) -> list[ExtractedUnit]:
    """Extract text from DOCX paragraphs and tables."""
    units = []
    try:
        doc = docx.Document(io.BytesIO(content))
    except Exception as e:
        raise ValueError("Failed to parse DOCX document")

    total_chars = 0
    # Extract paragraphs
    for p in doc.paragraphs:
        norm_text = normalize_text(p.text)
        if norm_text:
            total_chars = _check_char_bound(total_chars, len(norm_text))
            units.append(
                ExtractedUnit(
                    index=len(units),
                    text=norm_text,
                    source_label="Paragraph",
                )
            )

    # Extract tables
    for table in doc.tables:
        for row in table.rows:
            row_text = []
            for cell in row.cells:
                ct = normalize_text(cell.text)
                if ct:
                    row_text.append(ct)
            if row_text:
                combined = " | ".join(row_text)
                total_chars = _check_char_bound(total_chars, len(combined))
                units.append(
                    ExtractedUnit(
                        index=len(units),
                        text=combined,
                        source_label="Table Row",
                    )
                )

    if not units:
        raise ValueError("No extractable text found in DOCX")

    return units


def _extract_txt(content: bytes) -> list[ExtractedUnit]:
    """Extract text from TXT/MD, handling BOM and encoding safely."""
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError:
        try:
            text = content.decode("latin-1")
        except Exception:
            raise ValueError("Invalid text encoding, expected UTF-8")

    norm_text = normalize_text(text)
    if not norm_text:
        raise ValueError("Empty text document")
        
    _check_char_bound(0, len(norm_text))

    return [ExtractedUnit(index=0, text=norm_text, source_label="Document")]


def _extract_document_bytes(content: bytes, file_type: str) -> list[ExtractedUnit]:
    """Route extraction to the correct parser based on file type."""
    ft = file_type.lower()
    if ft == ".pdf":
        return _extract_pdf(content)
    elif ft == ".docx":
        return _extract_docx(content)
    elif ft in [".txt", ".md"]:
        return _extract_txt(content)
    else:
        raise ValueError(f"Unsupported file type for extraction: {ft}")


def process_document(
    user_id: UUID, workspace_id: UUID, document_id: UUID
) -> ExtractedDocument:
    """Retrieve from storage, extract text, and return structured result."""
    # 1. Verify ownership via DB
    doc = get_document(user_id, workspace_id, document_id)
    client = get_admin_client()

    # 2. Get storage path and update status to processing
    db_res = (
        client.table("documents")
        .select("storage_path")
        .eq("id", str(document_id))
        .execute()
    )
    if not db_res.data:
        raise HTTPException(status_code=404, detail="Document not found")
    storage_path = db_res.data[0]["storage_path"]

    # Mark as processing
    client.table("documents").update({"status": "processing"}).eq(
        "id", str(document_id)
    ).execute()

    # 3. Download from private storage
    try:
        storage_res = client.storage.from_(BUCKET_NAME).download(storage_path)
    except Exception as e:
        client.table("documents").update({"status": "failed"}).eq(
            "id", str(document_id)
        ).execute()
        raise HTTPException(
            status_code=status.HTTP_502_BAD_GATEWAY,
            detail="Failed to download document from storage",
        )

    # 4. Extract text safely
    try:
        units = _extract_document_bytes(storage_res, doc.file_type)
    except ValueError as e:
        client.table("documents").update({"status": "failed"}).eq(
            "id", str(document_id)
        ).execute()
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_CONTENT, detail=str(e)
        )
    except Exception as e:
        client.table("documents").update({"status": "failed"}).eq(
            "id", str(document_id)
        ).execute()
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail="Unexpected error during text extraction",
        )

    # 5. Mark as processed
    client.table("documents").update({"status": "processed"}).eq(
        "id", str(document_id)
    ).execute()

    char_count = sum(len(u.text) for u in units)

    return ExtractedDocument(
        document_id=doc.id,
        original_filename=doc.original_filename,
        file_type=doc.file_type,
        unit_count=len(units),
        character_count=char_count,
        units=units,
    )
