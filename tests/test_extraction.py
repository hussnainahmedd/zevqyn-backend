"""Unit tests for document extraction."""

from __future__ import annotations

import io
from unittest.mock import MagicMock, patch
import uuid

import docx
import pymupdf
import pytest
from fastapi.testclient import TestClient

from app.main import app

client = TestClient(app)

FAKE_USER_ID = str(uuid.uuid4())
FAKE_WORKSPACE_ID = str(uuid.uuid4())
FAKE_DOC_ID = str(uuid.uuid4())
FAKE_TOKEN = "valid-token"


@pytest.fixture(autouse=True)
def setup_auth(monkeypatch):
    """Automatically mock auth for all tests here unless overridden."""
    with patch("app.core.auth.get_admin_client") as mock_get:
        user = MagicMock()
        user.id = FAKE_USER_ID
        user.email = "test@example.com"
        response = MagicMock()
        response.user = user
        mock_get.return_value.auth.get_user.return_value = response
        yield mock_get


def _generate_test_pdf() -> bytes:
    """Generate a minimal multi-page PDF in memory."""
    doc = pymupdf.open()
    
    # Page 1
    page1 = doc.new_page()
    page1.insert_text((50, 50), "Hello from Page 1")
    
    # Page 2 (empty)
    doc.new_page()
    
    # Page 3
    page3 = doc.new_page()
    page3.insert_text((50, 50), "Hello from Page 3")
    
    pdf_bytes = doc.write()
    doc.close()
    return pdf_bytes


def _generate_test_docx() -> bytes:
    """Generate a minimal DOCX in memory."""
    doc = docx.Document()
    doc.add_paragraph("Paragraph 1 text.")
    doc.add_paragraph("Paragraph 2 text.")
    
    table = doc.add_table(rows=1, cols=2)
    row = table.rows[0]
    row.cells[0].text = "Cell A"
    row.cells[1].text = "Cell B"
    
    bio = io.BytesIO()
    doc.save(bio)
    return bio.getvalue()


@patch("app.services.extraction.get_document")
@patch("app.services.extraction.get_admin_client")
def test_extract_pdf(mock_db, mock_get_doc):
    # Mock doc info
    doc_meta = MagicMock()
    doc_meta.id = FAKE_DOC_ID
    doc_meta.original_filename = "test.pdf"
    doc_meta.file_type = ".pdf"
    mock_get_doc.return_value = doc_meta
    
    # Mock DB storage path retrieval
    mock_db_res = MagicMock()
    mock_db_res.data = [{"storage_path": "path/to/test.pdf"}]
    mock_db.return_value.table.return_value.select.return_value.eq.return_value.execute.return_value = mock_db_res
    
    # Mock storage download
    pdf_bytes = _generate_test_pdf()
    mock_db.return_value.storage.from_.return_value.download.return_value = pdf_bytes
    
    response = client.post(
        f"/api/v1/workspaces/{FAKE_WORKSPACE_ID}/documents/{FAKE_DOC_ID}/extract",
        headers={"Authorization": f"Bearer {FAKE_TOKEN}"}
    )
    
    assert response.status_code == 200
    data = response.json()
    assert data["file_type"] == ".pdf"
    assert data["unit_count"] == 2  # Page 1 and Page 3 (Page 2 was empty, should be skipped)
    
    units = data["units"]
    assert units[0]["page_number"] == 1
    assert "Page 1" in units[0]["text"]
    assert units[1]["page_number"] == 3
    assert "Page 3" in units[1]["text"]


@patch("app.services.extraction.get_document")
@patch("app.services.extraction.get_admin_client")
def test_extract_docx(mock_db, mock_get_doc):
    doc_meta = MagicMock()
    doc_meta.id = FAKE_DOC_ID
    doc_meta.original_filename = "test.docx"
    doc_meta.file_type = ".docx"
    mock_get_doc.return_value = doc_meta
    
    mock_db_res = MagicMock()
    mock_db_res.data = [{"storage_path": "path/to/test.docx"}]
    mock_db.return_value.table.return_value.select.return_value.eq.return_value.execute.return_value = mock_db_res
    
    docx_bytes = _generate_test_docx()
    mock_db.return_value.storage.from_.return_value.download.return_value = docx_bytes
    
    response = client.post(
        f"/api/v1/workspaces/{FAKE_WORKSPACE_ID}/documents/{FAKE_DOC_ID}/extract",
        headers={"Authorization": f"Bearer {FAKE_TOKEN}"}
    )
    
    assert response.status_code == 200
    data = response.json()
    assert data["file_type"] == ".docx"
    assert data["unit_count"] == 3  # P1, P2, Table Row
    
    units = data["units"]
    assert units[0]["page_number"] is None
    assert "Paragraph 1" in units[0]["text"]
    assert "Paragraph 2" in units[1]["text"]
    assert "Cell A | Cell B" in units[2]["text"]


@patch("app.services.extraction.get_document")
@patch("app.services.extraction.get_admin_client")
def test_extract_txt_utf8_bom(mock_db, mock_get_doc):
    doc_meta = MagicMock()
    doc_meta.id = FAKE_DOC_ID
    doc_meta.original_filename = "test.txt"
    doc_meta.file_type = ".txt"
    mock_get_doc.return_value = doc_meta
    
    mock_db_res = MagicMock()
    mock_db_res.data = [{"storage_path": "path/to/test.txt"}]
    mock_db.return_value.table.return_value.select.return_value.eq.return_value.execute.return_value = mock_db_res
    
    # Text with BOM and Windows line endings
    txt_bytes = b"\xef\xbb\xbfLine 1\r\nLine 2 \r\n\r\n\r\nLine 3"
    mock_db.return_value.storage.from_.return_value.download.return_value = txt_bytes
    
    response = client.post(
        f"/api/v1/workspaces/{FAKE_WORKSPACE_ID}/documents/{FAKE_DOC_ID}/extract",
        headers={"Authorization": f"Bearer {FAKE_TOKEN}"}
    )
    
    assert response.status_code == 200
    data = response.json()
    assert data["unit_count"] == 1
    
    text = data["units"][0]["text"]
    assert "Line 1\nLine 2\n\nLine 3" == text  # Normalized


@patch("app.services.extraction.get_document")
@patch("app.services.extraction.get_admin_client")
def test_extract_unsupported(mock_db, mock_get_doc):
    doc_meta = MagicMock()
    doc_meta.id = FAKE_DOC_ID
    doc_meta.original_filename = "test.exe"
    doc_meta.file_type = ".exe"
    mock_get_doc.return_value = doc_meta
    
    mock_db_res = MagicMock()
    mock_db_res.data = [{"storage_path": "path/to/test.exe"}]
    mock_db.return_value.table.return_value.select.return_value.eq.return_value.execute.return_value = mock_db_res
    
    mock_db.return_value.storage.from_.return_value.download.return_value = b"bad"
    
    response = client.post(
        f"/api/v1/workspaces/{FAKE_WORKSPACE_ID}/documents/{FAKE_DOC_ID}/extract",
        headers={"Authorization": f"Bearer {FAKE_TOKEN}"}
    )
    
    assert response.status_code == 422
    assert "Unsupported file type" in response.json()["detail"]


@patch("app.services.extraction.get_document")
@patch("app.services.extraction.get_admin_client")
def test_extract_storage_failure(mock_db, mock_get_doc):
    doc_meta = MagicMock()
    doc_meta.id = FAKE_DOC_ID
    doc_meta.original_filename = "test.txt"
    doc_meta.file_type = ".txt"
    mock_get_doc.return_value = doc_meta
    
    mock_db_res = MagicMock()
    mock_db_res.data = [{"storage_path": "path/to/test.txt"}]
    mock_db.return_value.table.return_value.select.return_value.eq.return_value.execute.return_value = mock_db_res
    
    mock_db.return_value.storage.from_.return_value.download.side_effect = Exception("Storage offline")
    
    response = client.post(
        f"/api/v1/workspaces/{FAKE_WORKSPACE_ID}/documents/{FAKE_DOC_ID}/extract",
        headers={"Authorization": f"Bearer {FAKE_TOKEN}"}
    )
    
    assert response.status_code == 502
    assert "download" in response.json()["detail"].lower()
